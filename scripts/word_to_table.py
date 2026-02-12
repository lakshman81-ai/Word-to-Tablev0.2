import pandas as pd
from docx import Document
from docx.document import Document as _Document
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph
import re
import argparse
import os
import json
import logging

# Get logger (it will inherit the handler from app_logging if setup)
logger = logging.getLogger(__name__)

# --- Helper Functions ---

def iter_block_items(parent):
    """
    Yield each paragraph and table child within *parent*, in document order.
    Each returned value is an instance of either Table or Paragraph.
    """
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("something's not right")

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)

def clean_text(text):
    if not text:
        return ""
    return re.sub(r'\s+', ' ', text).strip()

def get_table_header_text(table):
    """Extracts the first row text as a single string for matching."""
    if not table.rows:
        return ""
    
    # Try getting header from first row cells
    header_cells = table.rows[0].cells
    header_text = " ".join([clean_text(cell.text) for cell in header_cells])
    return header_text

def get_first_data_row_text(table):
    """Extracts the second row text (data row) for checks."""
    if len(table.rows) < 2:
        return ""
    data_cells = table.rows[1].cells
    return " ".join([clean_text(cell.text) for cell in data_cells])

def truncate_assemblies_context(text):
    """
    Truncates context for Assemblies.
    Rule: Truncated till 'DN' from string like "Drain or vent point DN 50 Fig.H DN 50 600"
    """
    match = re.search(r'^(.*?)\s+DN\b', text, re.IGNORECASE)
    if match:
        logger.debug(f"GATE: Truncating Assemblies Context. Pattern: 'DN'. Input: '{text}' -> Output: '{match.group(1).strip()}'")
        return match.group(1).strip()
    return text.strip()

def load_rules(rules_path="scripts/rules.json"):
    # Try multiple paths
    paths = [
        rules_path,
        os.path.join(os.path.dirname(__file__), "rules.json"),
        os.path.join(os.getcwd(), rules_path)
    ]

    for p in paths:
        if os.path.exists(p):
            try:
                with open(p, 'r') as f:
                    logger.info(f"Loading rules from {p}")
                    return json.load(f)
            except Exception as e:
                logger.error(f"Error loading rules from {p}: {e}")

    logger.warning("rules.json not found.")
    return []

def check_condition(cond, header_text, first_data_row, last_context):
    c_type = cond.get("type")
    val = cond.get("value", "")

    if c_type == "or":
        res = any(check_condition(sub, header_text, first_data_row, last_context) for sub in cond.get("conditions", []))
        logger.debug(f"GATE: OR Condition. Result: {res}")
        return res

    if c_type == "and":
        res = all(check_condition(sub, header_text, first_data_row, last_context) for sub in cond.get("conditions", []))
        logger.debug(f"GATE: AND Condition. Result: {res}")
        return res

    # Header checks
    if c_type == "header_contains":
        match = val in header_text
        if match: logger.debug(f"GATE: header_contains MATCHED. Val: '{val}' found in '{header_text[:30]}...'")
        return match
    if c_type == "header_contains_ignore_case":
        match = val.lower() in header_text.lower()
        if match: logger.debug(f"GATE: header_contains_ignore_case MATCHED. Val: '{val}' found in '{header_text[:30]}...'")
        return match
    if c_type == "not_header_contains_ignore_case":
        match = val.lower() not in header_text.lower()
        if match: logger.debug(f"GATE: not_header_contains_ignore_case MATCHED. Val: '{val}' NOT found in '{header_text[:30]}...'")
        return match
    if c_type == "header_regex":
        match = bool(re.search(val, header_text))
        if match: logger.debug(f"GATE: header_regex MATCHED. Regex: '{val}' matched '{header_text[:30]}...'")
        return match

    # Data checks
    if c_type == "data_contains":
        match = val in first_data_row
        if match: logger.debug(f"GATE: data_contains MATCHED. Val: '{val}' found in '{first_data_row[:30]}...'")
        return match
    if c_type == "data_regex":
        match = bool(re.search(val, first_data_row))
        if match: logger.debug(f"GATE: data_regex MATCHED. Regex: '{val}' matched '{first_data_row[:30]}...'")
        return match

    # Context checks
    if c_type == "context_contains":
        match = val in last_context
        if match: logger.debug(f"GATE: context_contains MATCHED. Val: '{val}' found in '{last_context[:30]}...'")
        return match
    if c_type == "context_contains_ignore_case":
        match = val.lower() in last_context.lower()
        if match: logger.debug(f"GATE: context_contains_ignore_case MATCHED. Val: '{val}' found in '{last_context[:30]}...'")
        return match
    if c_type == "not_context_contains_ignore_case":
        match = val.lower() not in last_context.lower()
        if match: logger.debug(f"GATE: not_context_contains_ignore_case MATCHED. Val: '{val}' NOT found in '{last_context[:30]}...'")
        return match

    return False

def evaluate_rule(rule, header_text, first_data_row, last_context):
    rule_id = rule.get("id", "Unknown")

    # Check main conditions
    if all(check_condition(c, header_text, first_data_row, last_context) for c in rule.get("conditions", [])):
        logger.info(f"Rule [{rule_id}] main conditions PASSED.")

        # Check subrules if any
        if "subrules" in rule:
            logger.debug(f"Rule [{rule_id}] checking subrules...")
            for sub in rule["subrules"]:
                if all(check_condition(c, header_text, first_data_row, last_context) for c in sub.get("conditions", [])):
                    target = sub.get("target")
                    logger.info(f"Rule [{rule_id}] SUBRULE MATCHED -> Target: {target}")
                    return target, rule.get("extra_cols", {})

            # If no subrule matches, check if main rule has a target
            if rule.get("target") != "VARIOUS":
                target = rule.get("target")
                logger.info(f"Rule [{rule_id}] no subrule matched, using Main Target: {target}")
                return target, rule.get("extra_cols", {})

            logger.warning(f"Rule [{rule_id}] matched but no valid subrule or target found.")
            return None, None

        target = rule.get("target")
        logger.info(f"Rule [{rule_id}] MATCHED -> Target: {target}")
        return target, rule.get("extra_cols", {})

    # logger.trace(f"Rule [{rule_id}] conditions failed.") # Only log if trace level or verbose debug
    return None, None

def get_extraction_logic_description():
    """Returns a string describing the current extraction rules."""
    rules = load_rules()
    if not rules:
        return "No rules loaded."

    desc = "Dynamic Extraction Logic:\n\n"
    for r in rules:
        desc += f"{r.get('id')}: Targets {r.get('target')}\n"
        if "description" in r:
            desc += f"  - {r['description']}\n"
    return desc

def parse_docx(docx_path, robust_mode=False):
    logger.info(f"Starting Parse Docx: {docx_path} (Robust: {robust_mode})")
    document = Document(docx_path)
    rules = load_rules()
    
    # Data storage
    data_store = {
        "Bolting_Data": [],
        "Assemblies_Data": [],
        "Reducing_Piping_Data": [],
        "Schedule_Table": [],
        "Branch_Conn": [],
        "Code_Exp": [],
        "Pipe_Table": [],
        "Flange_Table": [],
        "Fittings_Table": [],
        "Red_Fit_Table": [],
        "Valves_Table": [],
        "Instruments_Table": [],
        "Misc_Table": [],
        "Piping_Comps_Pg4": []
    }
    
    last_context = ""
    table_count = 0
    
    for block in iter_block_items(document):
        if isinstance(block, Paragraph):
            text = clean_text(block.text)
            if text:
                last_context = text
                # logger.debug(f"Context updated: '{text[:50]}...'")
        
        elif isinstance(block, Table):
            table_count += 1
            header_text = get_table_header_text(block)
            first_data_row = get_first_data_row_text(block)
            
            logger.debug(f"Analyzing Table #{table_count}. Header: '{header_text[:50]}...'")

            matched = False
            for rule in rules:
                target, extra_cols = evaluate_rule(rule, header_text, first_data_row, last_context)
                if target:
                    logger.info(f"Match Confirmed: Rule {rule.get('id')} -> Target {target}")
                    df = parse_table_to_df(block, robust_mode=robust_mode)

                    # Apply context actions
                    if rule.get("context_action") == "full":
                        logger.debug(f"Applying full context: '{last_context[:30]}...'")
                        df['Bolting Type'] = last_context
                    elif rule.get("context_action") == "truncated":
                        trunc = truncate_assemblies_context(last_context)
                        logger.debug(f"Applying truncated context: '{trunc[:30]}...'")
                        df['Assembly Type'] = trunc

                    # Apply extra columns
                    if extra_cols:
                        for k, v in extra_cols.items():
                            df[k] = v

                    # Ensure target list exists
                    if target not in data_store:
                        data_store[target] = []

                    data_store[target].append(df)

                    matched = True
                    break # Stop after first match
            
            if not matched:
                logger.debug(f"Table #{table_count} did not match any rules.")

    logger.info(f"Parse complete. Found {table_count} tables total.")
    return data_store

def parse_table_robust(table):
    """
    Parses a docx table handling merged cells by creating a grid map.
    """
    try:
        rows = len(table.rows)
        cols = len(table.columns)
        if rows == 0 or cols == 0:
            return pd.DataFrame()

        grid = [['' for _ in range(cols)] for _ in range(rows)]

        for r in range(rows):
            for c in range(cols):
                try:
                    # table.cell(r, c) handles merged cells automatically
                    cell = table.cell(r, c)
                    grid[r][c] = clean_text(cell.text)
                except IndexError:
                    continue

        # Use first row as headers
        headers = grid[0]

        # Deduplicate headers
        seen = {}
        final_headers = []
        for i, h in enumerate(headers):
            name = h if h else f"Column_{i+1}"
            if name in seen:
                seen[name] += 1
                name = f"{name}.{seen[name]}"
            else:
                seen[name] = 0
            final_headers.append(name)

        # Data rows
        data = []
        for r in range(1, rows):
            row_dict = {}
            for c in range(cols):
                if c < len(final_headers):
                    row_dict[final_headers[c]] = grid[r][c]
            data.append(row_dict)

        return pd.DataFrame(data)
    except Exception as e:
        logger.warning(f"Robust parsing failed: {e}. Falling back to simple parser.")
        return parse_table_to_df(table) # Fallback

def parse_table_to_df(table, robust_mode=False):
    """Converts a docx table to a pandas DataFrame."""
    if robust_mode:
        return parse_table_robust(table)

    data = []
    keys = None
    for i, row in enumerate(table.rows):
        text = [clean_text(cell.text) for cell in row.cells]
        if i == 0:
            keys = text
            continue
        # Handle cases where row length matches keys
        if keys and len(text) == len(keys):
            data.append(dict(zip(keys, text)))
        elif keys:
            # Fallback for merged cells or mismatches
            mapped = {}
            # Zip stops at shortest
            for k, v in zip(keys, text):
                mapped[k] = v
            data.append(mapped)
            
    return pd.DataFrame(data)

def save_data(data_store, output_dir):
    """Saves extracted data to CSVs."""
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)
    
    saved_files = []

    for key, dfs in data_store.items():
        if dfs:
            # Concatenate all tables of this type
            try:
                # Align columns before concat
                df_final = pd.concat(dfs, ignore_index=True, sort=False)
                path = os.path.join(output_dir, f"{key}.csv")
                df_final.to_csv(path, index=False)
                logger.info(f"Saved {key} to {path}")
                saved_files.append(path)
            except Exception as e:
                logger.error(f"Error saving {key}: {e}")
    
    return saved_files

if __name__ == "__main__":
    # Basic console logging setup if run directly
    logging.basicConfig(level=logging.DEBUG)

    parser = argparse.ArgumentParser(description="Extract specific tables from a Word Document")
    parser.add_argument("docx_path", help="Path to the .docx file")
    parser.add_argument("--output", default="output_tables", help="Output directory for CSVs")
    
    args = parser.parse_args()
    
    if not os.path.exists(args.docx_path):
        print(f"Error: File not found: {args.docx_path}")
        exit(1)
        
    # Simple CLI doesn't toggle robust mode yet, defaulting to False
    extracted_data = parse_docx(args.docx_path, robust_mode=False)
    save_data(extracted_data, args.output)
